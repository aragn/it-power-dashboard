"""
Fetch the EU ETS1 primary auction results (EU auctions on EEX, the common
auction platform) and write app/data/eua.json.

Sources (free, no credentials; the EEX DataSource API needs a licence):
  - current year:  EEX's yearly auction report
        https://public.eex-group.com/eex/eua-auction-report/
        emission-spot-primary-market-auction-report-<year>-data.xlsx
  - earlier years: EEX's archive of the yearly reports (2012-2025, zip)

Only successful auctions named "... EU" are kept (not the German, Polish or
Northern Irish ones).  The clearing price applies from the auction time
(11:00 CET) until the next auction; the dashboard draws it as a step line.
"""

import argparse
import io
import os
import zipfile
from datetime import date, datetime

import requests
from openpyxl import load_workbook

import compact

OUTPUT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "app",
    "data",
    "eua.json",
)

YEAR_URL = (
    "https://public.eex-group.com/eex/eua-auction-report/"
    "emission-spot-primary-market-auction-report-{year}-data.xlsx"
)
ARCHIVE_URL = (
    "https://www.eex.com/fileadmin/EEX/Downloads/Markets/Environmentals/"
    "EUA_Emission_Spot_Primary_Market_Auction_Report/Archive_Reports/"
    "emission-spot-primary-market-auction-report-2012-2025-data.zip"
)

FIRST_YEAR = 2025
HEADERS = {"User-Agent": "it-power-dashboard (github.com/aragn/it-power-dashboard)"}


def download(url):
    response = requests.get(url, headers=HEADERS, timeout=120)
    response.raise_for_status()
    return response.content


def parse_report(content):
    """Successful EU auctions of one yearly report."""
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    auctions = []
    for row in sheet.iter_rows(values_only=True):
        if not row or len(row) < 7 or not isinstance(row[1], datetime):
            continue
        name, status, price = str(row[3] or ""), str(row[5] or ""), row[6]
        if not name.strip().endswith(" EU") or status.lower() != "successful" or price is None:
            continue
        auction_time = row[2]
        auctions.append({
            "date": row[1].date().isoformat(),
            "time": auction_time.strftime("%H:%M") if isinstance(auction_time, datetime) else "11:00",
            "price": round(float(price), 2),
            "volume": float(row[11]) if len(row) > 11 and row[11] is not None else None,
        })
    return auctions


def year_report(year, archive):
    """The report of one year: the public yearly file, else the archive."""
    try:
        return download(YEAR_URL.format(year=year))
    except requests.HTTPError:
        name = next((n for n in archive.namelist() if n.endswith(f"-{year}-data.xlsx")), None)
        if name is None:
            raise
        return archive.read(name)


def main():
    parser = argparse.ArgumentParser(description="Fetch EU ETS1 EU auction clearing prices from EEX.")
    parser.parse_args()

    archive = zipfile.ZipFile(io.BytesIO(download(ARCHIVE_URL)))
    auctions = []
    for year in range(FIRST_YEAR, date.today().year + 1):
        found = parse_report(year_report(year, archive))
        print(f"{year}: {len(found)} EU auctions")
        auctions.extend(found)

    auctions.sort(key=lambda a: (a["date"], a["time"]))
    if not auctions:
        raise RuntimeError("No EU auctions found; not writing output.")

    compact.dump({
        "source": "EEX - EU ETS1 primary market auctions (EU)",
        "description": (
            "Clearing prices of the EU ETS1 EU auctions on EEX, EUR/tCO2, with "
            "auction date and time (CET/CEST) and volume (tCO2). A price "
            "applies from its auction until the next one."
        ),
        "auctions": auctions,
    }, OUTPUT_PATH)
    print(f"Wrote {len(auctions)} auctions, last {auctions[-1]['date']}: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
